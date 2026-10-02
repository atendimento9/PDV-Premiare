# -*- coding: utf-8 -*-
"""
Importador do catalogo Premiare Criativa.

Le a planilha oficial (fonte unica), extrai valores e imagens embutidas,
sanitiza texto por clausula e emite os JSONs publicos por ALLOWLIST.

Reproduzivel e idempotente. Nao escreve na planilha original.
Uso: python scripts/import-catalog.py
"""
from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "src" / "data"
DOCS_DIR = ROOT / "docs"
RAW_IMG_DIR = ROOT / ".cache" / "images-raw"

_env_dir = os.environ.get("PREMIARE_XLSX_DIR")
CANDIDATE_DIRS = [
    Path(_env_dir) if _env_dir else None,
    ROOT / "source",
    Path.home() / "Downloads",
]

HEADER_ROW = 4
FIRST_DATA_ROW = 5
LAST_DATA_ROW = 44

EXPECTED = {"products": 40, "categories": 9, "featured": 10, "kits": 5}

COL = {
    "id": 1, "category": 2, "name": 3, "sku": 4, "photo": 5, "link": 6,
    "materials": 7, "colors": 8, "dimensions": 9, "personalization": 10,
    "minQty": 11, "leadTime": 12, "differentiators": 13, "indicatedFor": 14,
    "application": 15, "kitCombination": 16, "potential": 17, "notes": 18,
    "sourceUrl": 19, "photoUrl": 20,
}

# Campos de texto publico que passam pelo sanitizador de clausulas.
SANITIZED_FIELDS = [
    "materials", "colors", "dimensions", "personalization",
    "differentiators", "indicatedFor", "application", "kitCombination", "notes",
]


# ---------------------------------------------------------------- utilidades

def strip_accents(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", str(s))
                   if unicodedata.category(c) != "Mn")


def slugify(s: str) -> str:
    s = strip_accents(s).lower()
    s = re.sub(r"[^a-z0-9]+", "-", s)
    return re.sub(r"^-+|-+$", "", s)[:80]


STOPWORDS = {
    "de", "da", "do", "das", "dos", "em", "com", "para", "e", "a", "o", "as", "os",
    "no", "na", "nos", "nas", "um", "uma", "ao", "aos", "por",
}


def tokens(s: str) -> list:
    raw = re.split(r"[^a-z0-9]+", strip_accents(s).lower())
    return [t for t in raw if t and t not in STOPWORDS]


def significant(ts: list) -> list:
    """Descarta tokens curtos, numericos e unidades: ruido em comparacoes."""
    units = {"mm", "cm", "ml", "mah", "kg", "un", "g", "l"}
    return [t for t in ts if len(t) >= 3 and not t.isdigit() and t not in units]


def token_match(a: str, b: str) -> bool:
    """Igualdade tolerante a plural e genero, sem confundir raizes distintas.

    Um prefixo curto e comum demais para servir de criterio: "neoplex" e
    "neoprene" compartilham quatro letras e sao produtos diferentes. Aceita-se
    apenas quando um token e prefixo do outro (plural) ou quando divergem so na
    ultima letra (genero: fosca/fosco).
    """
    if a == b:
        return True
    if min(len(a), len(b)) < 4:
        return False
    if a.startswith(b) or b.startswith(a):
        return True
    return len(a) == len(b) and a[:-1] == b[:-1]


def overlap_score(needle: list, haystack: list) -> float:
    if not needle:
        return 0.0
    hit = sum(1 for t in needle if any(token_match(t, h) for h in haystack))
    return hit / len(needle)


def name_url_score(name: str, url: str) -> float:
    """Confere a associacao imagem<->linha contra a URL oficial da foto.

    Varios nomes sao compostos ("Kit Boas-Vindas - trofeu, caneca e sacochila"):
    o slug da URL reflete o titulo do catalogo, e o trecho apos o travessao e uma
    enumeracao de conteudo escrita pela curadoria. Pontua-se o nome inteiro e
    tambem o titulo (antes do travessao), ficando com o melhor dos dois.
    """
    url_tokens = significant(tokens(Path(url).stem))
    if not url_tokens:
        return 0.0
    head = re.split(r"\s[–—-]\s|,", name)[0]
    return max(overlap_score(significant(tokens(cand)), url_tokens)
               for cand in (name, head))


# ------------------------------------------------- ausencia / sanitizacao

NI_EXACT = re.compile(r"^\s*ni\b|^\s*n[aa~o]*o\s+informado\s+no\s+site\s*$", re.I)

# Clausulas descartadas do texto publico. Motivos:
#  - referencia a fonte de terceiro ("o site", "Astor"): nao pode ir ao publico
#  - declaracao de ausencia: proibido exibir "nao informado"
#  - nota interna de curadoria/QA: nao e especificacao de produto
DROP_CLAUSE = [
    (re.compile(r"\bsites?\b", re.I), "referencia-a-fonte"),
    (re.compile(r"\bastor\b", re.I), "referencia-a-fornecedor"),
    (re.compile(r"\bNI\b"), "marcador-ausencia"),
    (re.compile(r"n\w*o\s+(foram\s+|fora\s+|\w{1,3}\s+)?"
                r"(informad|public|detalhad|listad|enumerad|divulgad|especificad|limitad|aplic)",
                re.I), "declaracao-de-ausencia"),
    (re.compile(r"\bsem\s+(nomes?|lista|enumerar|nomenclatura)", re.I), "declaracao-de-ausencia"),
    (re.compile(r"\bausentes?\b|\btruncad", re.I), "declaracao-de-ausencia"),
    (re.compile(r"^\s*(validar|solicitar|pedir|exigir|buscar|criar|montar|confirmar)\b", re.I),
     "nota-interna"),
    (re.compile(r"\bsele\w+o\b|\bcuradoria\b|parece\s+incomum|mantid\w+\s+exatamente", re.I),
     "nota-interna"),
]

CLAUSE_SPLIT = re.compile(r";\s*|(?<=\.)\s+(?=[A-ZÀ-Ú])")


def sanitize(value, dropped: list, ctx: str):
    """Remove, clausula a clausula, o que nao pode ser publicado.

    Retorna None quando nada sobra (o campo entao nao e renderizado).
    """
    if value is None:
        return None
    text = str(value).strip()
    if not text or NI_EXACT.match(text):
        return None
    kept = []
    for clause in CLAUSE_SPLIT.split(text):
        clause = clause.strip().strip(";").strip()
        if not clause:
            continue
        reason = None
        for rx, why in DROP_CLAUSE:
            if rx.search(clause):
                reason = why
                break
        # Comentario sobre o que a fonte publicou ou deixou de publicar, sem
        # dado nenhum junto ("Lista de funcoes publicada"). Uma clausula com
        # numero ("Area de gravacao publicada: 7 x 5 cm") carrega spec e fica.
        if reason is None and re.search(r"publicad", clause, re.I) \
                and not re.search(r"\d", clause):
            reason = "comentario-sobre-a-fonte"
        if reason:
            dropped.append({"campo": ctx, "clausula": clause, "motivo": reason})
        else:
            kept.append(clause.rstrip(" ."))
    if not kept:
        return None
    out = "; ".join(kept)
    return out[0].upper() + out[1:] + "."


MIN_QTY = re.compile(r"([\d.\s]+)\s*(unidades?|pe\w*as?|un\b)", re.I)


def parse_min_qty(value):
    if value is None:
        return None
    text = str(value).strip()
    if NI_EXACT.match(text):
        return None
    m = MIN_QTY.search(text)
    if not m:
        return None
    digits = re.sub(r"[^\d]", "", m.group(1))
    if not digits:
        return None
    unit = m.group(2).lower()
    unit = "unidades" if unit.startswith("un") else unit
    return {"value": int(digits), "unit": unit, "originalText": text}


# ------------------------------------------------------------ localizacao

def find_workbook() -> Path:
    pattern = re.compile(r"^relatorio_premiare_astor.*\.xlsx$", re.I)
    found = []
    for d in CANDIDATE_DIRS:
        if not d or not d.is_dir():
            continue
        for f in sorted(d.iterdir()):
            if pattern.match(f.name) and not f.name.startswith("~$"):
                found.append(f)
    if not found:
        sys.exit("BLOQUEIO 0.4.1: planilha 'relatorio_premiare_astor*.xlsx' nao encontrada.")

    def key(p: Path):
        m = re.search(r"\((\d+)\)", p.name)
        return (int(m.group(1)) if m else 0, p.name)

    return sorted(found, key=key)[-1]


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# ----------------------------------------------------------------- import

def main() -> int:
    xlsx = find_workbook()
    print("[fonte] %s" % xlsx)
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb["Produtos"]

    dropped = []
    issues = []

    # --- imagens embutidas, mapeadas por ancora de celula (nunca por ordem)
    by_row = {}
    for img in getattr(ws, "_images", []):
        frm = getattr(img.anchor, "_from", None)
        if frm is None:
            issues.append("Imagem sem ancora de celula - ignorada.")
            continue
        by_row.setdefault(frm.row + 1, []).append(img)

    # Guarda estrutural contra o erro grave e silencioso: um deslocamento de linha
    # associaria a foto errada a quase todos os produtos. Exigimos exatamente uma
    # imagem ancorada por linha de dado, sem sobra e sem falta.
    data_rows = set(range(FIRST_DATA_ROW, LAST_DATA_ROW + 1))
    anchored_rows = set(by_row)
    if anchored_rows != data_rows or any(len(v) != 1 for v in by_row.values()):
        issues.append(
            "Ancoragem de imagens fora do padrao esperado: linhas sem imagem %s; "
            "imagens fora da faixa de dados %s; linhas com mais de uma imagem %s."
            % (sorted(data_rows - anchored_rows), sorted(anchored_rows - data_rows),
               sorted(r for r, v in by_row.items() if len(v) != 1)))

    RAW_IMG_DIR.mkdir(parents=True, exist_ok=True)

    products = []
    internal = []
    media_rows = []

    for row in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
        def cell(key):
            return ws.cell(row=row, column=COL[key]).value

        name = str(cell("name") or "").strip()
        sku = str(cell("sku") or "").strip()
        category = str(cell("category") or "").strip()
        if not name or not sku or not category:
            issues.append("Linha %d: nome, codigo ou categoria ausente - produto descartado." % row)
            continue

        fields = {}
        for key in SANITIZED_FIELDS:
            fields[key] = sanitize(cell(key), dropped, "%s/%s" % (sku, key))

        lead = cell("leadTime")
        lead_time = None
        if lead is not None and not NI_EXACT.match(str(lead).strip()):
            lead_time = str(lead).strip()

        # --- imagem: ancora obrigatoria + verificacao cruzada com a URL oficial
        photo_url = str(cell("photoUrl") or "").strip()
        anchored = by_row.get(row, [])
        image_status = "sem_imagem"
        image_file = None
        xcheck = 0.0
        if len(anchored) > 1:
            image_status = "imagem_nao_confirmada"
            issues.append("Linha %d (%s): %d imagens ancoradas na mesma linha - imagem nao publicada."
                          % (row, sku, len(anchored)))
        elif len(anchored) == 1:
            img = anchored[0]
            xcheck = name_url_score(name, photo_url) if photo_url else 0.0
            if photo_url and xcheck < 0.5:
                image_status = "imagem_nao_confirmada"
                issues.append("Linha %d (%s): a URL oficial da foto nao confere com o nome do "
                              "produto (score %.2f) - imagem nao publicada." % (row, sku, xcheck))
            else:
                ext = str(getattr(img, "format", None) or "png").lower()
                image_file = RAW_IMG_DIR / ("%s.%s" % (slugify(sku), ext))
                image_file.write_bytes(img._data())
                image_status = "ok" if photo_url else "ok_sem_url_para_conferir"
        else:
            issues.append("Linha %d (%s): nenhuma imagem ancorada." % (row, sku))

        product = {
            "sku": sku,
            "slug": slugify(name),
            "name": name,
            "category": {"name": category, "slug": slugify(category)},
            "image": None,
            "materials": fields["materials"],
            "colors": fields["colors"],
            "dimensionsAndCapacity": fields["dimensions"],
            "personalization": fields["personalization"],
            "minimumQuantity": parse_min_qty(cell("minQty")),
            "leadTime": lead_time,
            "differentiators": fields["differentiators"],
            "indicatedFor": fields["indicatedFor"],
            "premiareApplication": fields["application"],
            "kitCombination": fields["kitCombination"],
            "technicalNotes": fields["notes"],
            "featured": False,
        }
        products.append(product)
        internal.append({
            "sku": sku, "row": row, "internalId": cell("id"),
            "potential": cell("potential"), "sourceUrl": cell("sourceUrl"),
            "photoUrl": photo_url, "rawNotes": cell("notes"),
            "imageStatus": image_status, "imageCrossCheck": round(xcheck, 3),
            "rawFile": image_file.name if image_file else None,
        })
        media_rows.append({
            "codigo": sku,
            "nome": name,
            "origem": "imagem embutida na planilha (ancora de celula)",
            "arquivo_bruto": image_file.name if image_file else "",
            "status_direitos": "autorizado pela Premiare (RIGHTS_CLEARANCE.md)",
            "status_imagem": image_status,
            "conferencia_url": "%.2f" % xcheck,
        })

    # --- slugs unicos, desempate estavel por SKU
    counts_slug = {}
    for p in products:
        counts_slug[p["slug"]] = counts_slug.get(p["slug"], 0) + 1
    dupes = set(s for s, n in counts_slug.items() if n > 1)
    for p in products:
        if p["slug"] in dupes:
            p["slug"] = "%s-%s" % (p["slug"], slugify(p["sku"]))
            issues.append("Slug duplicado desempatado por SKU: %s" % p["slug"])

    # ------------------------------------------------------------ Top 10
    ws_top = wb["Top 10"]
    by_sku = dict((p["sku"], p) for p in products)
    featured_order = []
    for r in range(5, ws_top.max_row + 1):
        sku = ws_top.cell(row=r, column=3).value
        pos = ws_top.cell(row=r, column=1).value
        if not sku:
            continue
        sku = str(sku).strip()
        if sku not in by_sku:
            issues.append("Top 10: SKU %s (posicao %s) nao encontrado na aba Produtos." % (sku, pos))
            continue
        by_sku[sku]["featured"] = True
        # ordem editorial: usada apenas para ordenar; a POSICAO nunca e publicada
        featured_order.append(sku)

    # ------------------------------------------------------- Kits sugeridos
    ws_kits = wb["Kits sugeridos"]
    kits = []
    for r in range(4, ws_kits.max_row + 1):
        kname = ws_kits.cell(row=r, column=1).value
        if not kname:
            continue
        kname = str(kname).strip()
        composition_raw = str(ws_kits.cell(row=r, column=3).value or "")
        parts = [c.strip() for c in re.split(r"\s\+\s|\s+ou\s+", composition_raw) if c.strip()]
        components = []
        for part in parts:
            pt = significant(tokens(part))
            scored = sorted(((overlap_score(pt, significant(tokens(p["name"]))), p)
                             for p in products), key=lambda x: -x[0])
            best_score, best = scored[0]
            # Empate no topo e ambiguidade: nao se adivinha, publica-se como texto.
            if len(scored) > 1 and scored[1][0] == best_score:
                issues.append('Kit "%s": componente "%s" casa com mais de um produto '
                              "(empate em %.2f) - publicado como texto." % (kname, part, best_score))
                best = None
            if best and best_score >= 0.6:
                components.append({"label": part, "slug": best["slug"], "sku": best["sku"]})
            else:
                components.append({"label": part, "slug": None, "sku": None})
                issues.append('Kit "%s": componente "%s" sem produto correspondente '
                              "(score %.2f) - publicado como texto." % (kname, part, best_score))
        kits.append({
            "slug": slugify(kname),
            "name": kname,
            "audience": str(ws_kits.cell(row=r, column=2).value or "").strip() or None,
            "concept": str(ws_kits.cell(row=r, column=4).value or "").strip() or None,
            "presentation": sanitize(ws_kits.cell(row=r, column=5).value, dropped, "kit/%s" % kname),
            "components": components,
        })

    # ---------------------------------------------------------- categorias
    order = []
    ws_res = wb["Resumo"]
    for r in range(8, 17):
        v = ws_res.cell(row=r, column=1).value
        if v:
            order.append(str(v).strip())
    cats = []
    for cname in order:
        items = [p for p in products if p["category"]["name"] == cname]
        if not items:
            issues.append("Categoria sem produto: %s - nao publicada." % cname)
            continue
        cover = next((p for p in items if p["featured"]), items[0])
        cats.append({"slug": slugify(cname), "name": cname,
                     "count": len(items), "coverSlug": cover["slug"]})
    for p in products:
        if p["category"]["name"] not in order:
            issues.append("Categoria fora do Resumo: %s (%s)." % (p["category"]["name"], p["sku"]))

    # ------------------------------------------------- imagens -> caminhos
    for p, meta in zip(products, internal):
        if meta["imageStatus"].startswith("ok"):
            base = "/catalog/products/%s/cover" % slugify(p["sku"])
            p["image"] = {
                "src": base + "-800.webp",
                "srcSmall": base + "-480.webp",
                "srcLarge": base + "-1200.webp",
                # alt DERIVADO do dado, nunca descricao visual
                "alt": "%s - %s" % (p["name"], p["category"]["name"]),
                "width": 800,
                "height": 800,
            }

    # ------------------------------------------------------ indice de busca
    search_index = []
    for p in products:
        hay = " ".join([x for x in [
            p["name"], p["sku"], p["category"]["name"], p["materials"], p["colors"],
            p["personalization"], p["differentiators"], p["indicatedFor"],
        ] if x])
        search_index.append({
            "slug": p["slug"], "name": p["name"], "sku": p["sku"],
            "category": p["category"]["name"], "categorySlug": p["category"]["slug"],
            "featured": p["featured"],
            "minQty": (p["minimumQuantity"] or {}).get("value"),
            "haystack": strip_accents(hay).lower(),
        })

    # --------------------------------------------------------------- saida
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    (ROOT / ".cache").mkdir(parents=True, exist_ok=True)

    def write_json(path, obj):
        path.write_text(json.dumps(obj, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    write_json(DATA_DIR / "catalog.generated.json", products)
    write_json(DATA_DIR / "categories.generated.json", cats)
    write_json(DATA_DIR / "featured.generated.json", featured_order)
    write_json(DATA_DIR / "kits.generated.json", kits)
    write_json(DATA_DIR / "catalog-search-index.json", search_index)
    write_json(ROOT / ".cache" / "internal.json", internal)

    with (DOCS_DIR / "MEDIA_MANIFEST.csv").open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(media_rows[0].keys()))
        w.writeheader()
        w.writerows(media_rows)

    counts = {
        "products": len(products), "categories": len(cats),
        "featured": len(featured_order), "kits": len(kits),
        "images_ok": sum(1 for m in internal if m["imageStatus"].startswith("ok")),
    }

    lines = [
        "# Relatorio de importacao do catalogo", "",
        "- Planilha: `%s`" % xlsx.name,
        "- SHA-256: `%s`" % sha256(xlsx),
        "- Aba `Produtos`: cabecalho na linha %d, dados %d-%d" % (HEADER_ROW, FIRST_DATA_ROW, LAST_DATA_ROW),
        "",
        "## Contagens (real vs. esperado)", "",
        "| Item | Real | Esperado | Situacao |", "|---|---|---|---|",
    ]
    for k in EXPECTED:
        real = counts[k]
        lines.append("| %s | %d | %d | %s |" % (k, real, EXPECTED[k],
                                                "OK" if real == EXPECTED[k] else "DIVERGENTE"))
    lines.append("| imagens publicaveis | %d | %d | %s |" % (
        counts["images_ok"], EXPECTED["products"],
        "OK" if counts["images_ok"] == EXPECTED["products"] else "DIVERGENTE"))
    lines += [
        "", "## Mapeamento de imagens", "",
        "Cada imagem foi associada ao produto pela **ancora de celula** do XLSX "
        "(coluna E, linhas 5-44), nunca pela ordem de aparicao. Cada associacao foi conferida "
        "contra o nome do arquivo da coluna `Foto oficial (URL)`; score abaixo de 0,50 marca o "
        "produto como `imagem_nao_confirmada` e a imagem nao e publicada.",
        "", "## Clausulas removidas do texto publico", "",
        "Total: %d. Regra: remove-se a clausula, nunca o campo inteiro, quando ela referencia a "
        "fonte de terceiro, declara ausencia de dado ou e nota interna de curadoria." % len(dropped),
        "", "| Campo | Clausula | Motivo |", "|---|---|---|",
    ]
    for d in dropped:
        lines.append("| `%s` | %s | %s |" % (d["campo"], d["clausula"].replace("|", "/"), d["motivo"]))
    (DOCS_DIR / "CATALOG_IMPORT_REPORT.md").write_text("\n".join(lines) + "\n", encoding="utf-8")

    idoc = ["# Ocorrencias do catalogo", "",
            "Registro fiel do que a planilha nao permitiu resolver. Nenhum dado foi ajustado "
            "para satisfazer contagem esperada (P1).", ""]
    idoc += ["- %s" % i for i in issues] if issues else ["- Nenhuma ocorrencia."]
    (DOCS_DIR / "CATALOG_ISSUES.md").write_text("\n".join(idoc) + "\n", encoding="utf-8")

    print(json.dumps(counts, indent=2))
    print("[clausulas removidas] %d" % len(dropped))
    print("[ocorrencias] %d" % len(issues))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
